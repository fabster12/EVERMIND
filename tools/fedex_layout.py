"""FedEx-branded Word layout helpers for the FDE capstone documents.

One module, one job: give every capstone .docx the same look so the
instructor sees a consistent house style at the Phase 1 approval gate.

Brand colours: FedEx purple #4D148C, FedEx orange #FF6600.
"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PURPLE = RGBColor(0x4D, 0x14, 0x8C)
ORANGE = RGBColor(0xFF, 0x66, 0x00)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT_GREY = RGBColor(0x76, 0x76, 0x76)
PURPLE_HEX = "4D148C"
ORANGE_HEX = "FF6600"
TINT_HEX = "F2EEF7"      # very light purple, for callout boxes
RULE_HEX = "D9D2E4"      # table rule


# --------------------------------------------------------------------------
# low-level xml helpers
# --------------------------------------------------------------------------

def shade_cell(cell, hex_fill):
    """Paint a solid background on a table cell."""
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def paragraph_border(paragraph, edge="bottom", hex_color=ORANGE_HEX, size=12, space=4):
    """Draw a single rule on one edge of a paragraph."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), hex_color)
    pBdr.append(el)


def _field(paragraph, instruction):
    """Insert a Word field code (used for PAGE / NUMPAGES)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    return run


def table_borders(table, hex_color=RULE_HEX, size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), hex_color)
        borders.append(el)
    tbl_pr.append(borders)


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:keepNext"))


def repeat_header(row):
    """Mark a table row as a header that repeats across page breaks."""
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


# --------------------------------------------------------------------------
# document shell
# --------------------------------------------------------------------------

def new_document(footer_left="FedEx FDE Capstone"):
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name in ("List Bullet", "List Number"):
        st = doc.styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(10.5)
        st.paragraph_format.space_after = Pt(2)

    _build_footer(section, footer_left)
    return doc


def _build_footer(section, footer_left):
    footer = section.footer
    p = footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    paragraph_border(p, edge="top", hex_color=RULE_HEX, size=4, space=6)

    p.add_run(footer_left + "\t")
    p.add_run("Page ")
    _field(p, "PAGE")
    p.add_run(" of ")
    _field(p, "NUMPAGES")
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = LIGHT_GREY


# --------------------------------------------------------------------------
# content blocks
# --------------------------------------------------------------------------

def banner(doc, title, subtitle=None, kicker="FedEx  |  FDE Training Program"):
    """Purple title block with the orange FedEx rule underneath."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, PURPLE_HEX)
    cell.width = Cm(16.6)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(kicker.upper())
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(10 if not subtitle else 2)
    r2 = p2.add_run(title)
    r2.bold = True
    r2.font.size = Pt(22)
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if subtitle:
        p3 = cell.add_paragraph()
        p3.paragraph_format.space_before = Pt(0)
        p3.paragraph_format.space_after = Pt(10)
        r3 = p3.add_run(subtitle)
        r3.font.size = Pt(10.5)
        r3.font.color.rgb = RGBColor(0xEB, 0xE2, 0xF5)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(12)
    paragraph_border(rule, edge="bottom", hex_color=ORANGE_HEX, size=18, space=0)
    return table


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    keep_with_next(p)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(15)
    r.font.color.rgb = PURPLE
    paragraph_border(p, edge="bottom", hex_color=ORANGE_HEX, size=8, space=4)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    keep_with_next(p)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = PURPLE
    return p


def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    keep_with_next(p)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x2B, 0x2B, 0x2B)
    return p


def labelled(doc, label, text, label_color=PURPLE):
    """'Problem. <text>' - the four-question format in one paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    lr = p.add_run(label.rstrip(".") + ". ")
    lr.bold = True
    lr.font.color.rgb = label_color
    p.add_run(text)
    return p


def body(doc, text, italic=False, color=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if size is not None:
        r.font.size = Pt(size)
    return p


def hint(doc, text):
    """Grey italic guidance line - used in the blank template only."""
    return body(doc, text, italic=True, color=GREY, size=9.5)


def bullets(doc, items, style="List Bullet"):
    out = []
    for item in items:
        p = doc.add_paragraph(style=style)
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(lead)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)
        out.append(p)
    return out


def callout(doc, title, text):
    """Light purple box for a note the reviewer should not miss."""
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, TINT_HEX)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = PURPLE
    r.font.size = Pt(10)
    t = p.add_run(text)
    t.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def grid(doc, headers, rows, widths=None, font_size=9.5):
    """Simple header-banded table."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_borders(table)
    repeat_header(table.rows[0])

    for idx, head in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, PURPLE_HEX)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(head)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_vals in rows:
        cells = table.add_row().cells
        for idx, val in enumerate(row_vals):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.size = Pt(font_size)

    if widths:
        for row in table.rows:
            for idx, w in enumerate(widths):
                row.cells[idx].width = Cm(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def meta_block(doc, pairs):
    """Two-column 'Group / Lead / Date' strip under the banner."""
    table = doc.add_table(rows=0, cols=2)
    for key, val in pairs:
        cells = table.add_row().cells
        cells[0].width = Cm(4.0)
        cells[1].width = Cm(12.6)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(key)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = PURPLE
        p2 = cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(1)
        r2 = p2.add_run(val)
        r2.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def page_break(doc):
    doc.add_page_break()
